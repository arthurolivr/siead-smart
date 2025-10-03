import os
from openpyxl import Workbook, load_workbook

class ExcelLogger:
    def __init__(self, filename):
        self.filename = filename
        if not os.path.exists(filename):
            wb = Workbook()
            ws = wb.active
            ws.append(["Pergunta"])
            wb.save(filename)

    def log_question(self, question):
        wb = load_workbook(self.filename)
        ws = wb.active
        ws.append([question])
        wb.save(self.filename)